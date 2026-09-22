"""
Excel report generator for rate intelligence output.
Produces a formatted .xlsx with:
  - Summary tab: key metrics at a glance
  - All Rates tab: full interpreted rate data
  - Target Property tab: The Tyrwhitt rates across channels
  - Competitor Analysis tab: competitor ADR comparison
  - Anomalies tab: flagged records
"""

import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter
from datetime import datetime
from collections import defaultdict
from typing import Optional
import os


# ── Colour palette ──────────────────────────────────────────────────────────
DARK_NAVY   = "0D1B2A"
GOLD        = "C9A84C"
TEAL        = "2A9D8F"
LIGHT_GREY  = "F5F5F5"
WHITE       = "FFFFFF"
RED         = "C0392B"
GREEN       = "27AE60"
AMBER       = "E67E22"

def _fill(hex_color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=hex_color)

def _font(bold=False, color=WHITE, size=10, name="Calibri") -> Font:
    return Font(bold=bold, color=color, size=size, name=name)

def _border() -> Border:
    thin = Side(style="thin", color="CCCCCC")
    return Border(left=thin, right=thin, top=thin, bottom=thin)

def _center() -> Alignment:
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def _left() -> Alignment:
    return Alignment(horizontal="left", vertical="center", wrap_text=True)


RATE_TYPE_COLORS = {
    "NR":      "E74C3C",  # red
    "FLEX":    "27AE60",  # green
    "BB":      "2980B9",  # blue
    "DBB":     "8E44AD",  # purple
    "SPA":     "16A085",  # teal
    "RO":      "7F8C8D",  # grey
    "UNKNOWN": "BDC3C7",  # light grey
}

TIER_COLORS = {
    "MORE_EXPENSIVE": "E74C3C",
    "COMPARABLE":     "E67E22",
    "CHEAPER":        "27AE60",
    "TARGET":         C9A84C if (C9A84C := "C9A84C") else "C9A84C",
    "UNKNOWN":        "BDC3C7",
}

CHANNEL_LABELS = {
    "booking_com":  "Booking.com",
    "expedia":      "Expedia",
    "agoda":        "Agoda",
    "lekkeslaap":   "LekkeSlaap",
    "sa_venues":    "SA Venues",
}


def _header_row(ws, headers: list[str], row=1, fill_hex=DARK_NAVY):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.fill = _fill(fill_hex)
        cell.font = _font(bold=True, color=WHITE)
        cell.alignment = _center()
        cell.border = _border()


def _autofit(ws, min_width=10, max_width=40):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(min_width, min(max_len + 2, max_width))


# ── Tab builders ──────────────────────────────────────────────────────────────

def _build_summary_tab(wb, interpreted: list[dict], target_name: str, run_date: str):
    ws = wb.create_sheet("Summary")
    ws.sheet_view.showGridLines = False

    # Title block
    ws.merge_cells("A1:F1")
    title_cell = ws["A1"]
    title_cell.value = f"Rate Intelligence Report — {target_name}"
    title_cell.fill = _fill(DARK_NAVY)
    title_cell.font = Font(bold=True, color=GOLD, size=14, name="Calibri")
    title_cell.alignment = _center()
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:F2")
    sub = ws["A2"]
    sub.value = f"Report generated: {run_date}  |  Data: OTA scrape + LLM interpretation"
    sub.fill = _fill(TEAL)
    sub.font = _font(bold=False, color=WHITE, size=9)
    sub.alignment = _center()

    # KPIs
    target_rates = [r for r in interpreted if r.get("is_target_property")]
    comp_rates    = [r for r in interpreted if not r.get("is_target_property") and r.get("price_zar")]
    all_prices    = [r["price_zar"] for r in interpreted if r.get("price_zar")]

    target_prices  = [r["price_zar"] for r in target_rates if r.get("price_zar")]
    target_avg     = round(sum(target_prices) / len(target_prices), 2) if target_prices else None

    comp_prices    = [r["price_zar"] for r in comp_rates]
    comp_avg       = round(sum(comp_prices) / len(comp_prices), 2) if comp_prices else None

    channels_seen  = len({r["channel"] for r in interpreted})
    total_records  = len(interpreted)
    anomaly_count  = sum(1 for r in interpreted if r.get("anomaly_flags"))

    kpis = [
        ("Total rate records", total_records),
        ("Channels scraped", channels_seen),
        (f"{target_name} avg rate (ZAR)", f"R {target_avg:,.0f}" if target_avg else "N/A"),
        ("Competitor avg rate (ZAR)", f"R {comp_avg:,.0f}" if comp_avg else "N/A"),
        ("Rate parity delta", (
            f"R {abs(target_avg - comp_avg):,.0f} {'above' if target_avg > comp_avg else 'below'} market"
            if target_avg and comp_avg else "N/A"
        )),
        ("Anomalies flagged", anomaly_count),
    ]

    ws.row_dimensions[3].height = 6
    for i, (label, value) in enumerate(kpis, start=4):
        label_cell = ws.cell(row=i, column=1, value=label)
        label_cell.fill = _fill(LIGHT_GREY)
        label_cell.font = _font(color=DARK_NAVY, bold=True, size=10)
        label_cell.alignment = _left()
        label_cell.border = _border()

        ws.merge_cells(f"B{i}:F{i}")
        val_cell = ws.cell(row=i, column=2, value=value)
        val_cell.font = _font(color=DARK_NAVY, size=10)
        val_cell.alignment = _left()
        val_cell.border = _border()
        ws.row_dimensions[i].height = 20

    # Channel breakdown
    ch_row = 4 + len(kpis) + 2
    ws.merge_cells(f"A{ch_row}:F{ch_row}")
    ws.cell(row=ch_row, column=1, value="Channel Breakdown").fill = _fill(TEAL)
    ws.cell(row=ch_row, column=1).font = _font(bold=True)
    ws.cell(row=ch_row, column=1).alignment = _center()

    _header_row(ws, ["Channel", "Records", "Avg Rate (ZAR)", "Min Rate", "Max Rate", "Anomalies"],
                row=ch_row + 1, fill_hex=DARK_NAVY)

    ch_data = defaultdict(list)
    for r in interpreted:
        ch_data[r["channel"]].append(r)

    for ci, (ch, recs) in enumerate(sorted(ch_data.items()), start=ch_row + 2):
        prices = [r["price_zar"] for r in recs if r.get("price_zar")]
        anom = sum(1 for r in recs if r.get("anomaly_flags"))
        row_data = [
            CHANNEL_LABELS.get(ch, ch),
            len(recs),
            f"R {sum(prices)/len(prices):,.0f}" if prices else "N/A",
            f"R {min(prices):,.0f}" if prices else "N/A",
            f"R {max(prices):,.0f}" if prices else "N/A",
            anom,
        ]
        bg = LIGHT_GREY if ci % 2 == 0 else WHITE
        for j, val in enumerate(row_data, 1):
            cell = ws.cell(row=ci, column=j, value=val)
            cell.fill = _fill(bg)
            cell.font = _font(color=DARK_NAVY, size=10)
            cell.alignment = _center()
            cell.border = _border()
        ws.row_dimensions[ci].height = 18

    _autofit(ws)
    ws.column_dimensions["A"].width = 30


def _build_all_rates_tab(wb, interpreted: list[dict]):
    ws = wb.create_sheet("All Rates")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"

    headers = [
        "Channel", "Property", "Check-in", "Check-out",
        "Room (raw)", "Rate Plan (raw)", "Price (ZAR)",
        "Rate Type", "Room Type", "Meals", "Cancellation",
        "Comp Tier", "Confidence", "Anomalies", "Notes"
    ]
    _header_row(ws, headers)
    ws.row_dimensions[1].height = 24

    for i, r in enumerate(interpreted, start=2):
        bg = LIGHT_GREY if i % 2 == 0 else WHITE
        rate_type = r.get("rate_type", "UNKNOWN")
        tier = r.get("competitor_tier", "UNKNOWN")
        anom = ", ".join(r.get("anomaly_flags", []))

        row_vals = [
            CHANNEL_LABELS.get(r.get("channel", ""), r.get("channel", "")),
            r.get("property_name", ""),
            r.get("check_in", ""),
            r.get("check_out", ""),
            r.get("room_label_raw", ""),
            r.get("rate_label_raw", ""),
            r.get("price_zar", ""),
            rate_type,
            r.get("room_type", ""),
            r.get("meals_included", ""),
            r.get("cancellation", ""),
            tier,
            r.get("confidence", ""),
            anom,
            r.get("notes", ""),
        ]

        for j, val in enumerate(row_vals, 1):
            cell = ws.cell(row=i, column=j, value=val)
            cell.fill = _fill(bg)
            cell.alignment = _left()
            cell.border = _border()
            cell.font = _font(color=DARK_NAVY, size=9)

            # Colour code rate type
            if j == 8 and rate_type in RATE_TYPE_COLORS:
                cell.fill = _fill(RATE_TYPE_COLORS[rate_type])
                cell.font = _font(color=WHITE, size=9)
            # Colour code tier
            if j == 12 and tier in TIER_COLORS:
                cell.fill = _fill(TIER_COLORS[tier])
                cell.font = _font(color=WHITE, size=9)
            # Format price
            if j == 7 and isinstance(val, (int, float)):
                cell.number_format = 'R #,##0.00'

        ws.row_dimensions[i].height = 18

    _autofit(ws)


def _build_target_tab(wb, interpreted: list[dict], target_name: str):
    ws = wb.create_sheet("Target Property")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:G1")
    t = ws["A1"]
    t.value = f"{target_name} — Rate Summary by Channel"
    t.fill = _fill(DARK_NAVY)
    t.font = Font(bold=True, color=GOLD, size=12)
    t.alignment = _center()
    ws.row_dimensions[1].height = 28

    target = [r for r in interpreted if r.get("is_target_property")]
    if not target:
        ws.cell(row=2, column=1, value="No target property rates found.")
        return

    _header_row(ws, ["Channel", "Room Type", "Rate Type", "Meals", "Cancellation", "Price (ZAR)", "Confidence"],
                row=2, fill_hex=TEAL)

    for i, r in enumerate(target, start=3):
        bg = LIGHT_GREY if i % 2 == 0 else WHITE
        rate_type = r.get("rate_type", "UNKNOWN")
        row_vals = [
            CHANNEL_LABELS.get(r.get("channel", ""), r.get("channel", "")),
            r.get("room_type", ""),
            rate_type,
            r.get("meals_included", ""),
            r.get("cancellation", ""),
            r.get("price_zar", ""),
            r.get("confidence", ""),
        ]
        for j, val in enumerate(row_vals, 1):
            cell = ws.cell(row=i, column=j, value=val)
            cell.fill = _fill(bg)
            cell.font = _font(color=DARK_NAVY, size=10)
            cell.alignment = _center()
            cell.border = _border()
            if j == 3 and rate_type in RATE_TYPE_COLORS:
                cell.fill = _fill(RATE_TYPE_COLORS[rate_type])
                cell.font = _font(color=WHITE, size=10)
            if j == 6 and isinstance(val, (int, float)):
                cell.number_format = 'R #,##0.00'
        ws.row_dimensions[i].height = 18

    _autofit(ws)


def _build_competitor_tab(wb, interpreted: list[dict]):
    ws = wb.create_sheet("Competitor Analysis")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:G1")
    t = ws["A1"]
    t.value = "Competitor Rate Benchmarking"
    t.fill = _fill(DARK_NAVY)
    t.font = Font(bold=True, color=GOLD, size=12)
    t.alignment = _center()
    ws.row_dimensions[1].height = 28

    comps = [r for r in interpreted if not r.get("is_target_property") and r.get("price_zar")]
    if not comps:
        ws.cell(row=2, column=1, value="No competitor rates found.")
        return

    _header_row(ws, ["Property", "Channel", "Room Type", "Rate Type", "Price (ZAR)", "vs Target", "Confidence"],
                row=2, fill_hex=TEAL)

    # Sort by price desc
    comps_sorted = sorted(comps, key=lambda x: x.get("price_zar", 0), reverse=True)

    for i, r in enumerate(comps_sorted, start=3):
        bg = LIGHT_GREY if i % 2 == 0 else WHITE
        tier = r.get("competitor_tier", "UNKNOWN")
        rate_type = r.get("rate_type", "UNKNOWN")
        row_vals = [
            r.get("property_name", ""),
            CHANNEL_LABELS.get(r.get("channel", ""), r.get("channel", "")),
            r.get("room_type", ""),
            rate_type,
            r.get("price_zar", ""),
            tier,
            r.get("confidence", ""),
        ]
        for j, val in enumerate(row_vals, 1):
            cell = ws.cell(row=i, column=j, value=val)
            cell.fill = _fill(bg)
            cell.font = _font(color=DARK_NAVY, size=10)
            cell.alignment = _center()
            cell.border = _border()
            if j == 6 and tier in TIER_COLORS:
                cell.fill = _fill(TIER_COLORS[tier])
                cell.font = _font(color=WHITE, size=10)
            if j == 5 and isinstance(val, (int, float)):
                cell.number_format = 'R #,##0.00'
        ws.row_dimensions[i].height = 18

    _autofit(ws)


def _build_anomalies_tab(wb, interpreted: list[dict]):
    ws = wb.create_sheet("Anomalies")
    ws.sheet_view.showGridLines = False

    flagged = [r for r in interpreted if r.get("anomaly_flags")]
    if not flagged:
        ws.cell(row=1, column=1, value="✓ No anomalies detected in this scrape run.")
        ws["A1"].font = _font(color=GREEN, bold=True)
        return

    _header_row(ws, ["Channel", "Property", "Room (raw)", "Price (ZAR)", "Flags", "Notes"],
                fill_hex=RED)

    for i, r in enumerate(flagged, start=2):
        bg = LIGHT_GREY if i % 2 == 0 else WHITE
        row_vals = [
            CHANNEL_LABELS.get(r.get("channel", ""), r.get("channel", "")),
            r.get("property_name", ""),
            r.get("room_label_raw", ""),
            r.get("price_zar", ""),
            ", ".join(r.get("anomaly_flags", [])),
            r.get("notes", ""),
        ]
        for j, val in enumerate(row_vals, 1):
            cell = ws.cell(row=i, column=j, value=val)
            cell.fill = _fill(bg)
            cell.font = _font(color=DARK_NAVY, size=10)
            cell.alignment = _left()
            cell.border = _border()
            if j == 4 and isinstance(val, (int, float)):
                cell.number_format = 'R #,##0.00'
        ws.row_dimensions[i].height = 18

    _autofit(ws)


# ── Public entry point ────────────────────────────────────────────────────────

def generate_report(
    interpreted: list[dict],
    output_path: str,
    target_name: str = "The Tyrwhitt",
) -> str:
    """
    Generate the Excel rate intelligence report.
    Returns the path to the saved file.
    """
    run_date = datetime.now().strftime("%Y-%m-%d %H:%M")

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    _build_summary_tab(wb, interpreted, target_name, run_date)
    _build_all_rates_tab(wb, interpreted)
    _build_target_tab(wb, interpreted, target_name)
    _build_competitor_tab(wb, interpreted)
    _build_anomalies_tab(wb, interpreted)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path
